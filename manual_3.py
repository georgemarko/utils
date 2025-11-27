import requests
import shutil
import zipfile
import os
import csv
from lxml import etree
from docx import Document
from credentials import MYUSERNAME, MYPASSWORD

############### AUTHENTICATION #######################

######################################################

############## INPUTS ################################

# If SELECT_ONE_IMO is True then only the VESSEL_IMO is used

SELECT_ONE_IMO = True
VESSEL_IMO = 9723019
COMPANY_NAME = "DYNACOM"

######################################################


# ---------------- CONFIG ----------------
API_BASE_URL = "https://mariner.alphamrn.com/api"
EMISSION_SOURCES_URL = f"{API_BASE_URL}/emission-sources"
CSV_FILE = "vessels.csv"
WORD_TEMPLATE = "model_3.docx"
TEMP_DIR = "temp_docx_extract"




def authenticate(username, password, base_url=API_BASE_URL, remember_me=False):
    url = f"{base_url}/authenticate"
    
    payload = {
        "username": username,
        "password": password,
        "rememberMe": remember_me
    }
    
    response = requests.post(url, json=payload)
    
    if response.status_code == 200:
        data = response.json()
        jwt_token = data.get("id_token")
        return jwt_token
    else:
        print(f"Authentication failed: {response.status_code}")
        print(response.text)
        return None

token = authenticate(MYUSERNAME, MYPASSWORD)

HEADERS = {
    "Authorization": f"Bearer {token}",
    "Accept": "application/json"
}



# ---------------- PLACEHOLDER MAPPING ----------------
PLACEHOLDER_MAP = {
    "{{VSLNAME}}": "vesselName",
    "{{IMO}}": "imo",
    "{{DWT}}": "deadWeight",
    "{{DWTVALUE}}": "deadWeightValue",
    "{{HULL}}": "hullNo",
    "{{VSLTYPE}}": "vesselType",
    "{{VSLTYPENAME}}": "vesselTypeName",
    "{{DWG}}": "dwg",  # From CSV
    "{{COUNTRY}}": "flagCountryName",
    "{{PORT}}": "registryPort",
    "{{CALLSIGN}}": "callsign",
    "{{GROSSTONNAGE}}": "grossTonnage",
    "{{NETTONNAGE}}": "netTonnage",
    "{{EEDI}}": "aEedi",
    "{{EEXI}}": "aEexi",
    "{{ICECLASS}}": "iceClass",
    "{{BUILDER}}": "shipbuilder",
    "{{YEAR}}": "deliveryDate",
    "{{LENGTHOA}}": "overallLength",
    "{{LENGTHBP}}": "lengthBp",
    "{{BREADTH}}": "breadth",
    "{{DEPTH}}": "depth",
    "{{SLD}}": "summerLoadDraught"
}

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# ---------------- DOCX UTILITIES ----------------
def replace_placeholders_in_t_nodes(t_nodes, placeholders):
    """
    Replace placeholders inside each <w:t> node individually.
    Handles multiple placeholders in the same run and surrounding static text.
    """
    for t in t_nodes:
        if t.text:
            for ph, replacement in placeholders.items():
                if ph in t.text:
                    t.text = t.text.replace(ph, str(replacement))
    return True


def recursive_replace(element, placeholders, visited=None):
    """
    Recursively replace all placeholders in Word elements,
    including tables, text boxes, content controls, headers, footers, etc.
    """
    if visited is None:
        visited = set()

    element_id = id(element)
    if element_id in visited:
        return False
    visited.add(element_id)

    replaced = False

    # Replace in current text nodes
    t_nodes = list(element.iter("{%s}t" % W_NS))
    if t_nodes:
        if replace_placeholders_in_t_nodes(t_nodes, placeholders):
            replaced = True

    # Recurse into shapes / text boxes
    for drawing in element.iter("{%s}drawing" % W_NS):
        for txbx in drawing.iter("{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}txbxContent"):
            if recursive_replace(txbx, placeholders, visited):
                replaced = True

    # Recurse into content controls (<w:sdtContent>)
    for sdt in element.iter("{%s}sdtContent" % W_NS):
        if sdt is not element:
            if recursive_replace(sdt, placeholders, visited):
                replaced = True

    return replaced



def process_docx(input_path, output_path, placeholders):
    if os.path.exists(TEMP_DIR):
        shutil.rmtree(TEMP_DIR)
    os.makedirs(TEMP_DIR)
    with zipfile.ZipFile(input_path, 'r') as zip_ref:
        zip_ref.extractall(TEMP_DIR)
    word_dir = os.path.join(TEMP_DIR, "word")
    for root_dir, dirs, files in os.walk(word_dir):
        for file in files:
            if file.endswith(".xml"):
                file_path = os.path.join(root_dir, file)
                parser = etree.XMLParser(remove_blank_text=False)
                tree = etree.parse(file_path, parser)
                root = tree.getroot()
                recursive_replace(root, placeholders)
                tree.write(file_path, encoding="UTF-8", xml_declaration=True)
    with zipfile.ZipFile(output_path, "w") as zipf:
        for root_dir, dirs, files in os.walk(TEMP_DIR):
            for file in files:
                file_path = os.path.join(root_dir, file)
                arcname = os.path.relpath(file_path, TEMP_DIR)
                zipf.write(file_path, arcname)
    shutil.rmtree(TEMP_DIR)

# ---------------- CSV UTILITIES ----------------
def get_imos_for_company(csv_file, company_name=None, vessel_imo=None):
    imos = {}
    with open(csv_file, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            if vessel_imo != None:
                if row["IMO"] == f"{vessel_imo}":
                    imos[row["IMO"]] = row
            elif company_name != None:
                if row["COMPANY NAME"] == company_name:
                    imos[row["IMO"]] = row
    return imos

# ---------------- API FETCH ----------------
def get_vessel(imo):
    resp = requests.get(f"{API_BASE_URL}/vessels/imo/{imo}", headers=HEADERS)
    resp.raise_for_status()
    return resp.json()

def format_number(value):
    """Helper to format numeric values with thousand separators"""
    try:
        return f"{float(value):,}"
    except (TypeError, ValueError):
        return "N/A"

# ---------------- PLACEHOLDER HELPERS ----------------
def format_vessel_placeholder(vessel, csv_dwg):
    result = {}
    for ph, field in PLACEHOLDER_MAP.items():
        if field == "dwg":
            value = csv_dwg
        elif field == "vesselTypeName":
            # UPPERCASE for {{VSLTYPENAME}}
            value = vessel.get("vesselType", "")
            value = " ".join(word.upper() for word in value.split('_'))
        elif field == "vesselType":
            # Capitalized for {{VSLTYPE}}
            value = vessel.get("vesselType", "")
            value = " ".join(word.capitalize() for word in value.split('_'))
        elif field == "aEedi":
            eedi = vessel.get("aEedi")
            value = f"{format_number(eedi)} gr CO2 / ton-mile" if eedi else "N/A"

        elif field == "aEexi":
            eexi = vessel.get("aEexi")
            value = f"{format_number(eexi)} gr CO2 / ton-mile" if eexi else "N/A"

        # Numbers with thousand separators
        elif field in ["grossTonnage", "netTonnage"]:
            num = vessel.get(field)
            if num is not None:
                value = f"{num:,.2f}".rstrip("0").rstrip(".")  # preserves decimals if any
            else:
                value = "N/A"
        elif field == "deadWeightValue":
            num = vessel.get("deadWeight")
            if num is not None:
                value = f"{num:,.2f}".rstrip("0").rstrip(".")
            else:
                value = "N/A"
        elif field == "deadWeight":
            num = vessel.get(field)
            if num is not None:
                val = f"{num:,.2f}".rstrip(".")  # preserves decimals if any
                value = f"{val} MT"
            else:
                value = "N/A"
        elif field in ["overallLength", "lengthBp", "breadth", "depth", "summerLoadDraught"]:
            num = vessel.get(field)
            if num is not None:
                val = f"{num:,.2f}".rstrip(".")
                value = f"{val} m"
            else:
                value = "N/A"
        elif "." in field:
            parts = field.split(".")
            val = vessel
            for part in parts:
                val = val.get(part, {})
            value = val if isinstance(val, str) else str(val)
        elif field == "deliveryDate":
            value = vessel.get("deliveryDate", "N/A")
            if value != "N/A" and value:
                value = value.split("T")[0].split("-")[0]
        else:
            value = vessel.get(field, "N/A")

        if value is None:
            value = "N/A"
        result[ph] = str(value)
    return result


def get_issue_number(doc):
    """
    Find the table with 'Issue Number' text and return the last cell value
    from that column.
    """
    for table in doc.tables:
        # Find the row and column index with "Issue Number"
        issue_col_idx = None
        for row in table.rows:
            for cell_idx, cell in enumerate(row.cells):
                cell_text = cell.text.strip()
                if "Issue Number" in cell_text or "Issue No" in cell_text:
                    issue_col_idx = cell_idx
                    break
            if issue_col_idx is not None:
                break
        
        # If found, get the last row's value in that column
        if issue_col_idx is not None:
            for row in reversed(table.rows):
                cell_value = row.cells[issue_col_idx].text.strip()
                if cell_value and cell_value != "Issue Number" and cell_value != "Issue No":
                    return cell_value
    
    print("Issue number not found, using default: 00")
    return "00"  # Default if not found


# ---------------- MAIN SCRIPT ----------------
if SELECT_ONE_IMO:
    imos = get_imos_for_company(CSV_FILE, vessel_imo=VESSEL_IMO)
else:
    imos = get_imos_for_company(CSV_FILE, company_name=COMPANY_NAME)
for imo, csv_row in imos.items():
    vessel = get_vessel(imo)
    print(f"Processing vessel with IMO {imo}")
    csv_dwg = csv_row.get("DWG NO.", "UNKNOWN")
    print(f"DWG is {csv_dwg}")
    placeholders = format_vessel_placeholder(vessel, csv_dwg)

    # Get issue number for filename
    doc = Document(WORD_TEMPLATE)
    issue_num = get_issue_number(doc)

    output_filename = f"{csv_dwg} {vessel['vesselName']} – SEEMP PART III Issue No. {issue_num}"
    output_doc = f"{output_filename}.docx"

    process_docx(WORD_TEMPLATE, output_doc, placeholders)
